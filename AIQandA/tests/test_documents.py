"""Tests for document parsing and chunking."""

from __future__ import annotations

import pytest
from django.core.files.base import ContentFile
from docx import Document as DocxDocument

from AIQandA.models import KnowledgeDocument
from AIQandA.services.documents import (
    DocumentIngestError,
    extract_pages,
    ingest_document,
    normalize_text,
    split_page_text,
)

pytestmark = pytest.mark.django_db


def test_normalize_text_removes_blank_lines() -> None:
    assert normalize_text("第一行\r\n\r\n  第二行  \n") == "第一行\n第二行"


def test_split_page_text_respects_limit_and_overlap() -> None:
    text = "\n".join(f"第{i}段内容" + "说明" * 60 for i in range(6))
    chunks = split_page_text(text, max_chars=200, overlap=30)
    assert len(chunks) > 1
    assert all(len(chunk) <= 200 for chunk in chunks)


def test_ingest_document_creates_ready_chunks(document) -> None:
    assert document.status == "ready"
    assert document.page_count == 1
    assert document.chunks.count() >= 1
    assert "监督学习" in document.extracted_text


def test_extract_rejects_unsupported_extension(course, teacher, tmp_path, settings) -> None:
    settings.MEDIA_ROOT = tmp_path
    record = KnowledgeDocument(course=course, uploaded_by=teacher, title="坏文件")
    record.file.save("bad.xyz", ContentFile(b"hello"), save=False)
    record.save()
    with pytest.raises(DocumentIngestError):
        extract_pages(record)


def test_extract_docx_paragraphs(course, teacher, tmp_path, settings) -> None:
    settings.MEDIA_ROOT = tmp_path
    source = tmp_path / "source.docx"
    docx = DocxDocument()
    docx.add_paragraph("支持向量机是一种监督学习方法。")
    docx.save(source)

    record = KnowledgeDocument(course=course, uploaded_by=teacher, title="DOCX 资料")
    record.file.save("source.docx", ContentFile(source.read_bytes()), save=False)
    record.save()
    pages = extract_pages(record)
    assert pages
    assert "支持向量机" in pages[0][1]


def test_failed_ingest_marks_document(course, teacher, tmp_path, settings) -> None:
    settings.MEDIA_ROOT = tmp_path
    record = KnowledgeDocument(course=course, uploaded_by=teacher, title="损坏 PDF")
    record.file.save("broken.pdf", ContentFile(b"not a real pdf"), save=False)
    record.save()
    with pytest.raises(DocumentIngestError):
        ingest_document(record)
    record.refresh_from_db()
    assert record.status == KnowledgeDocument.Status.FAILED
    assert record.error_message
