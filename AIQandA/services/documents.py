"""Document parsing and chunking for course materials."""

from __future__ import annotations

import logging
import unicodedata
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from docx import Document as DocxDocument
from pypdf import PdfReader

from ..models import DocumentChunk, KnowledgeDocument
from .embeddings import embed_texts

logger = logging.getLogger(__name__)

CHUNK_MAX_CHARS = 700
CHUNK_OVERLAP = 120
MAX_EXTRACTED_CHARS = 2_000_000


class DocumentIngestError(RuntimeError):
    """Raised when a document cannot be parsed or indexed."""


def normalize_text(text: str) -> str:
    """Normalize line endings and whitespace without destroying paragraphs."""
    normalized = unicodedata.normalize("NFKC", text)
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in normalized.split("\n")]
    compact = "\n".join(line for line in lines if line)
    return compact.strip()


def _extract_pdf(path: Path) -> list[tuple[int | None, str]]:
    reader = PdfReader(str(path))
    pages: list[tuple[int | None, str]] = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append((index, page.extract_text() or ""))
    return pages


def _extract_docx(path: Path) -> list[tuple[int | None, str]]:
    document = DocxDocument(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return [(None, "\n".join(parts))]


def _extract_text_file(path: Path) -> list[tuple[int | None, str]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return [(None, raw.decode(encoding))]
        except UnicodeDecodeError:
            continue
    raise DocumentIngestError("文本编码无法识别，请使用 UTF-8 或 GB18030 编码。")


def extract_pages(document: KnowledgeDocument) -> list[tuple[int | None, str]]:
    """Extract page-aware text from the uploaded file."""
    path = Path(document.file.path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = _extract_pdf(path)
    elif suffix == ".docx":
        pages = _extract_docx(path)
    elif suffix in {".txt", ".md"}:
        pages = _extract_text_file(path)
    else:
        raise DocumentIngestError(f"不支持的文件类型：{suffix}")
    normalized = [
        (page_number, normalize_text(text)) for page_number, text in pages if normalize_text(text)
    ]
    if not normalized:
        raise DocumentIngestError("未从文件中提取到可读文本。")
    return normalized


def split_page_text(
    text: str,
    max_chars: int = CHUNK_MAX_CHARS,
    overlap: int = CHUNK_OVERLAP,
) -> list[str]:
    """Split one page into overlapping, paragraph-aware chunks."""
    paragraphs = [paragraph.strip() for paragraph in text.split("\n") if paragraph.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current)
                current = ""
            start = 0
            while start < len(paragraph):
                end = min(start + max_chars, len(paragraph))
                chunks.append(paragraph[start:end])
                if end >= len(paragraph):
                    break
                start = end - overlap
            continue

        candidate = f"{current}\n{paragraph}" if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
            continue

        chunks.append(current)
        tail = current[-overlap:] if overlap and len(current) > overlap else current
        candidate = f"{tail}\n{paragraph}" if tail else paragraph
        current = candidate if len(candidate) <= max_chars else paragraph

    if current:
        chunks.append(current)
    return [chunk.strip() for chunk in chunks if chunk.strip()]


def _store_embeddings(chunks: list[DocumentChunk]) -> None:
    if not settings.AI_EMBEDDING_MODEL or not settings.AI_API_KEY:
        return
    vectors = embed_texts([chunk.content for chunk in chunks])
    if not vectors:
        return
    for chunk, vector in zip(chunks, vectors):
        chunk.embedding = vector
    DocumentChunk.objects.bulk_update(chunks, ["embedding"])


def ingest_document(document: KnowledgeDocument) -> int:
    """Parse, chunk, and index a document. Returns the chunk count."""
    try:
        pages = extract_pages(document)
        extracted = "\n\n".join(text for _, text in pages)
        if len(extracted) > MAX_EXTRACTED_CHARS:
            extracted = extracted[:MAX_EXTRACTED_CHARS]

        payload: list[DocumentChunk] = []
        index = 0
        for page_number, text in pages:
            for chunk_text in split_page_text(text):
                payload.append(
                    DocumentChunk(
                        document=document,
                        chunk_index=index,
                        page_number=page_number,
                        content=chunk_text,
                        char_count=len(chunk_text),
                    )
                )
                index += 1

        if not payload:
            raise DocumentIngestError("文档切块后没有可用内容。")

        with transaction.atomic():
            document.chunks.all().delete()
            DocumentChunk.objects.bulk_create(payload)
            document.extracted_text = extracted
            document.page_count = len(pages)
            document.status = KnowledgeDocument.Status.READY
            document.error_message = ""
            document.processed_at = timezone.now()
            document.save(
                update_fields=[
                    "extracted_text",
                    "page_count",
                    "status",
                    "error_message",
                    "processed_at",
                ]
            )
        chunks = list(document.chunks.order_by("chunk_index"))
        _store_embeddings(chunks)
        return len(chunks)
    except Exception as exc:  # noqa: BLE001 - persisted for the teacher to inspect
        if isinstance(exc, DocumentIngestError):
            message = str(exc)
        else:
            message = f"文档处理失败：{exc}"
        document.status = KnowledgeDocument.Status.FAILED
        document.error_message = message[:2000]
        document.processed_at = timezone.now()
        document.save(update_fields=["status", "error_message", "processed_at"])
        logger.exception("Failed to ingest document %s", document.pk)
        raise DocumentIngestError(message) from exc
